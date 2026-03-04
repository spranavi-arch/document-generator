"""
FastAPI endpoints for the document generator pipeline:
- GET /fields-and-values  — return required field names and current values
- GET /draft-preview      — return assembled draft text (after prepare + field values)
- POST /field-values      — submit/update field values (JSON body)
- POST /fetch-field-values — run CURL to call external API, obtain field values (one request per field)
- POST /prepare           — upload sample docs, run sectioning + extraction + prompts (sets fields)
"""
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from io import BytesIO

# Ensure project root and docgen are on path
_root = Path(__file__).resolve().parent.parent
_docgen = Path(__file__).resolve().parent
for p in (str(_root), str(_docgen)):
    if p not in sys.path:
        sys.path.insert(0, p)

from fastapi import FastAPI, UploadFile, File, HTTPException, Body

from docgen.sectioner import Sectioner
from docgen.extractor import Extractor
from docgen.section_prompt_generator import SectionPromptGenerator
from docgen.section_generator import SectionGenerator
from docgen.assembler import Assembler
from docgen.utils import fill_placeholders_from_context_with_llm, fill_placeholders_from_field_values
from docgen.field_fetcher import FieldFetcher
from docgen.question_generator import QuestionGenerator


# ---------------------------------------------------------------------------
# File to text: LibreOffice when SOFFICE_PATH set (preserves list numbering), else python-docx
# ---------------------------------------------------------------------------
def _html_to_text_with_numbering(soup) -> str:
    """Extract text from BeautifulSoup, preserving list numbering (1., 2., 3.) from <ol>/<li>."""
    from bs4 import NavigableString, Tag
    lines = []

    def visit(el):
        if isinstance(el, NavigableString):
            s = str(el).strip()
            if s:
                lines.append(s)
            return
        if not isinstance(el, Tag):
            return
        tag = el.name
        if tag == "ol":
            for i, li in enumerate(el.find_all("li", recursive=False), start=1):
                part = li.get_text(separator=" ", strip=True)
                if part:
                    lines.append(f"{i}. {part}")
            return
        if tag == "ul":
            for li in el.find_all("li", recursive=False):
                part = li.get_text(separator=" ", strip=True)
                if part:
                    lines.append(f"• {part}")
            return
        if tag == "li":
            part = el.get_text(separator=" ", strip=True)
            if part:
                lines.append(part)
            return
        for child in el.children:
            visit(child)

    body = soup.find("body") or soup
    for child in body.children:
        visit(child)
    return "\n".join(lines) if lines else soup.get_text(separator="\n")


def _docx_to_text_via_libreoffice(data: bytes, soffice_path: str) -> str | None:
    """Convert DOCX to text via LibreOffice (HTML). Returns None on failure."""
    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir = Path(tmpdir)
            input_path = tmpdir / "input.docx"
            input_path.write_bytes(data)
            subprocess.run(
                [
                    soffice_path,
                    "--headless",
                    "--convert-to",
                    "html",
                    "--outdir",
                    str(tmpdir),
                    str(input_path),
                ],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=60,
            )
            html_path = tmpdir / "input.html"
            if not html_path.exists():
                return None
            html = html_path.read_text(encoding="utf-8", errors="ignore")
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(html, "html.parser")
            text = _html_to_text_with_numbering(soup)
            if not text.strip():
                text = soup.get_text(separator="\n")
            return text.strip()
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, FileNotFoundError, OSError, Exception):
        return None


def _docx_to_text_python_docx(data: bytes) -> str:
    """Extract plain text from DOCX using python-docx (paragraphs only)."""
    from docx import Document
    doc = Document(BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs).strip()


def _file_to_text(data: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".docx"):
        try:
            soffice_path = None
            try:
                from docgen.config import SOFFICE_PATH as _sp
                soffice_path = _sp
            except Exception:
                pass
            if soffice_path and os.path.isfile(soffice_path):
                text = _docx_to_text_via_libreoffice(data, soffice_path)
                if text is not None:
                    return text
            return _docx_to_text_python_docx(data)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Invalid .docx: {e}") from e
    try:
        return data.decode("utf-8", errors="replace").strip()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not decode file: {e}") from e


# ---------------------------------------------------------------------------
# In-memory session state (single session)
# ---------------------------------------------------------------------------
_state: dict = {
    "blueprint": None,
    "sections_list": None,
    "extracted_samples": None,
    "section_prompts": None,
    "field_values": None,
    "required_fields": None,
}


def _ensure_prepared():
    if _state["blueprint"] is None or _state["section_prompts"] is None:
        raise HTTPException(
            status_code=409,
            detail="Call POST /prepare first with two sample documents to set fields.",
        )


# ---------------------------------------------------------------------------
# FastAPI app
# ---------------------------------------------------------------------------
app = FastAPI(
    title="Document Generator API",
    description="Fields, values, and draft preview for the docgen pipeline.",
)


@app.post("/prepare")
async def prepare(
    sample1: UploadFile = File(...),
    sample2: UploadFile = File(...),
):
    """
    Upload two sample documents (.txt or .docx). Runs sectioning, extraction,
    and prompt generation. After this, GET /fields-and-values and GET /draft-preview
    are available; use POST /field-values to set values before fetching the draft.
    """
    raw1 = await sample1.read()
    raw2 = await sample2.read()
    s1 = _file_to_text(raw1, sample1.filename or "")
    s2 = _file_to_text(raw2, sample2.filename or "")

    if not (s1 or s2):
        raise HTTPException(status_code=400, detail="Both documents are empty.")

    sectioner = Sectioner()
    extractor = Extractor()
    section_prompt_generator = SectionPromptGenerator()

    try:
        blueprint = sectioner.divide_into_sections(s1, s2)
    except ValueError as e:
        raise HTTPException(status_code=422, detail={"step": "sectioning", "message": str(e)}) from e

    sections_list = blueprint["sections"]
    try:
        extracted_samples = extractor.extract_sections_from_docs(s1, s2, sections_list)
    except Exception as e:
        raise HTTPException(
            status_code=422,
            detail={"step": "extraction", "message": str(e)},
        ) from e

    section_prompts = []
    for i, sec in enumerate(sections_list):
        sample_text = extracted_samples[i] if i < len(extracted_samples) else ""
        section_prompts.append(
            section_prompt_generator.generate_prompt_and_fields(
                sec["name"], sec.get("purpose", ""), sample_text
            )
        )

    required_fields = []
    seen = set()
    for info in section_prompts:
        for f in info.get("required_fields", []):
            if f and f not in seen:
                seen.add(f)
                required_fields.append(f)

    _state["blueprint"] = blueprint
    _state["sections_list"] = sections_list
    _state["extracted_samples"] = extracted_samples
    _state["section_prompts"] = section_prompts
    _state["required_fields"] = required_fields
    _state["field_values"] = {}

    return {
        "fields": required_fields,
        "values": {},
        "sections": [s["name"] for s in sections_list],
    }


@app.get("/fields-and-values")
async def get_fields_and_values():
    """
    Return the list of required field names and their current values.
    Call POST /prepare first to set fields; use POST /field-values to set values.
    """
    _ensure_prepared()
    values = _state.get("field_values") or {}
    fields = _state.get("required_fields") or []
    return {"fields": fields, "values": values}


@app.post("/field-values")
async def post_field_values(body: dict = Body(..., examples=[{"plaintiff_name": "Jane Doe", "case_number": "12345"}])):
    """
    Submit or update field values. Body is a JSON object with field names as keys
    and string values. Merges into existing values. Use before GET /draft-preview.
    """
    _ensure_prepared()
    new_values = {k: v if isinstance(v, str) else str(v) for k, v in (body or {}).items() if v is not None}
    current = _state.get("field_values") or {}
    current.update(new_values)
    _state["field_values"] = current
    return {"values": current}


@app.post("/fetch-field-values")
async def fetch_field_values(
    body: dict = Body(
        ...,
        examples=[
            {
                "curl": "curl 'https://api.example.com/chat' -X POST -H 'Authorization: Bearer TOKEN' -H 'Content-Type: application/json' --data-raw '{\"content\": \"REPLACE\"}'",
                "include_case_summary": True,
            }
        ],
    ),
):
    """
    Call an external chat API using the provided CURL command to obtain values for all required fields.
    The API is called once per field with a generated question; answers are extracted and stored.
    Requires POST /prepare first. Optionally fetches a case summary and stores it as case_summary.
    """
    _ensure_prepared()
    curl_str = (body or {}).get("curl") or (body or {}).get("curl_command") or ""
    curl_str = curl_str.strip() if isinstance(curl_str, str) else ""
    if not curl_str:
        raise HTTPException(
            status_code=400,
            detail="Body must include 'curl' or 'curl_command' with a CURL command string.",
        )
    include_case_summary = (body or {}).get("include_case_summary", True)

    required_fields = _state.get("required_fields") or []
    if not required_fields:
        current = _state.get("field_values") or {}
        return {"values": current, "message": "No required fields to fetch."}

    question_generator = QuestionGenerator()
    field_fetcher = FieldFetcher()

    field_to_question = question_generator.generate_questions_for_fields(required_fields)
    field_values = field_fetcher.fetch_all_fields_via_chat(
        curl_str, required_fields, field_to_question
    )
    if include_case_summary:
        case_summary = field_fetcher.fetch_case_summary(curl_str)
        field_values["case_summary"] = case_summary or ""

    current = _state.get("field_values") or {}
    current.update(field_values)
    _state["field_values"] = current

    return {"values": current, "fields_fetched": list(field_values.keys())}


@app.get("/draft-preview")
async def get_draft_preview():
    """
    Generate and return the assembled draft text using current field values.
    Call POST /prepare first, then POST /field-values (or rely on empty values),
    then this endpoint.
    """
    _ensure_prepared()
    blueprint = _state["blueprint"]
    sections_list = _state["sections_list"]
    section_prompts = _state["section_prompts"]
    extracted_samples = _state["extracted_samples"]
    field_values = _state.get("field_values") or {}

    section_generator = SectionGenerator()
    assembler = Assembler()

    section_texts_ordered = []
    for i, sec in enumerate(sections_list):
        name = sec["name"]
        info = section_prompts[i]
        prompt = info.get("prompt", "")
        required_fields = info.get("required_fields", [])
        section_field_values = {f: field_values.get(f, "") for f in required_fields}
        sample_text = extracted_samples[i] if i < len(extracted_samples) else ""
        text = section_generator.generate_section(
            prompt,
            section_field_values,
            sample_text=sample_text,
            section_name=name,
        )
        section_texts_ordered.append(text)

    final_draft = assembler.assemble(blueprint, section_texts_ordered)
    final_draft = fill_placeholders_from_context_with_llm(final_draft, field_values)
    final_draft = fill_placeholders_from_field_values(final_draft, field_values)

    return {"draft": final_draft or ""}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("docgen.api:app", host="0.0.0.0", port=8000, reload=True)
