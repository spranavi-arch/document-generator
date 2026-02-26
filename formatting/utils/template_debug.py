"""
Debug utilities for template validation before fill.
Run on template.docx to verify placeholders and run structure.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

try:
    from utils.template_filler import iter_body_blocks
    from utils.placeholder_docx import paragraph_full_text
except ImportError:
    from template_filler import iter_body_blocks
    from placeholder_docx import paragraph_full_text


def debug_placeholders(docx_path: str | Path) -> None:
    """
    Print every paragraph (and table-cell paragraph) containing `{{...}}`,
    plus its run breakdown. Use this to validate the stored template before fill.

    Look for:
    - Paragraphs like `1. {{PLAINTIFF_NAME}},{{COUNTY}}` (static text already gone?)
    - Duplicated markers like `{{LICENSE_PLATE}}{{LICENSE_PLATE}}`
    - `COUNTY OF {{COUNTY}}` vs only `{{COUNTY}}` (missing static prefix?)
    """
    print(debug_placeholders_report(docx_path))


def debug_placeholders_report(docx_path: str | Path) -> str:
    """Same as debug_placeholders but returns the report as a string (for UI). Run-safe: uses paragraph_full_text."""
    doc = Document(docx_path)
    pat = re.compile(r"\{\{[^}]+\}\}")
    lines: list[str] = []
    for i, (p, table_id, row, col) in enumerate(iter_body_blocks(doc)):
        full = paragraph_full_text(p)
        if not pat.search(full):
            continue
        loc = f"table[{table_id}].row[{row}].cell[{col}]" if table_id is not None else f"paragraph[{i}]"
        lines.append(f"\n[{loc}] {full}")
        for rj, run in enumerate(p.runs):
            t = (run.text or "").replace("\n", "\\n")
            if t.strip():
                lines.append(f"  run[{rj}]: {t}")
    return "\n".join(lines) if lines else "(No placeholders found.)"


if __name__ == "__main__":
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else "output/template.docx"
    debug_placeholders(path)
