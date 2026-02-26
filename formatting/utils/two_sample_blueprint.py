"""
Deterministic diff + LLM classification pipeline for learning placeholders from two sample DOCX files.

Data model: TextUnit (extracted in document order), SpanDiff (minimal changed spans with context).
Steps: extract units → align units → char-span diffs → LLM classify → apply placeholders to doc.
"""

from __future__ import annotations

import difflib
import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional

from docx import Document

try:
    from utils.style_extractor import iter_body_blocks
except ImportError:
    from style_extractor import iter_body_blocks


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------


@dataclass
class TextUnit:
    unit_id: str
    text: str
    raw_text: str
    kind: str  # "paragraph" or "table_cell"
    path: str
    para_ref: Any = None  # docx Paragraph for apply phase


@dataclass
class SpanDiff:
    unit_id_a: str
    unit_id_b: str
    unit_text_a: str
    unit_text_b: str
    left_context: str
    old_span: str
    new_span: str
    right_context: str
    char_start_a: int
    char_end_a: int
    char_start_b: int
    char_end_b: int
    anchor: str


# ---------------------------------------------------------------------------
# Step 1: Extract text units (document order = iter_body_blocks)
# ---------------------------------------------------------------------------


def normalize_text(s: str) -> str:
    s = (s or "").replace("\u00a0", " ")
    s = "\n".join(line.rstrip() for line in s.splitlines())
    s = " ".join(s.split())
    return s.strip()


def extract_units(doc: Document, docx_path: str = "") -> list[TextUnit]:
    """
    Extract paragraphs and table-cell paragraphs in document order (same as iter_body_blocks).
    unit_id is u:0000, u:0001, ... so index i maps to para_id i in template_builder.
    """
    units: list[TextUnit] = []
    for i, (para, table_id, row, col) in enumerate(iter_body_blocks(doc)):
        raw = (para.text or "").strip()
        text = normalize_text(raw)
        if table_id is not None:
            kind = "table_cell"
            path = f"table[{table_id}].row[{row}].cell[{col}]"
        else:
            kind = "paragraph"
            path = f"paragraph[{i}]"
        units.append(TextUnit(
            unit_id=f"u:{i:04d}",
            text=text,
            raw_text=raw,
            kind=kind,
            path=path,
            para_ref=para,
        ))
    return units


# ---------------------------------------------------------------------------
# Step 2: Align units between two docs
# ---------------------------------------------------------------------------


def align_units(
    units_a: list[TextUnit],
    units_b: list[TextUnit],
    similarity_threshold: float = 0.65,
) -> list[tuple[Optional[TextUnit], Optional[TextUnit]]]:
    """
    Returns list of aligned pairs (unit_a, unit_b). Uses SequenceMatcher on texts,
    then for replace/delete/insert blocks does local best-effort matching by similarity.
    """
    if not units_a and not units_b:
        return []
    a_texts = [u.text for u in units_a]
    b_texts = [u.text for u in units_b]
    sm = difflib.SequenceMatcher(a=a_texts, b=b_texts, autojunk=False)
    aligned: list[tuple[Optional[TextUnit], Optional[TextUnit]]] = []

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            for i, j in zip(range(i1, i2), range(j1, j2)):
                aligned.append((units_a[i], units_b[j]))
        elif tag in ("replace", "delete", "insert"):
            a_block = units_a[i1:i2]
            b_block = units_b[j1:j2]
            used_b = set()
            for ua in a_block:
                best = None
                best_score = 0.0
                for idx, ub in enumerate(b_block):
                    if idx in used_b:
                        continue
                    score = difflib.SequenceMatcher(
                        None, ua.text, ub.text, autojunk=False
                    ).ratio()
                    if score > best_score:
                        best_score = score
                        best = (idx, ub)
                if best and best_score >= similarity_threshold:
                    used_b.add(best[0])
                    aligned.append((ua, best[1]))
                else:
                    aligned.append((ua, None))
            for idx, ub in enumerate(b_block):
                if idx not in used_b:
                    aligned.append((None, ub))
    return aligned


# ---------------------------------------------------------------------------
# Step 3: Char-level span diffs inside aligned units
# ---------------------------------------------------------------------------


def char_span_diffs(
    text_a: str,
    text_b: str,
    context_chars: int = 30,
) -> list[tuple[int, int, int, int, str, str, str, str]]:
    """
    Returns list of (a0, a1, b0, b1, left, old, new, right).
    """
    sm = difflib.SequenceMatcher(None, text_a, text_b, autojunk=False)
    diffs: list[tuple[int, int, int, int, str, str, str, str]] = []

    for tag, a0, a1, b0, b1 in sm.get_opcodes():
        if tag == "equal":
            continue
        if tag in ("replace", "delete", "insert"):
            old = text_a[a0:a1]
            new = text_b[b0:b1]
            left = text_a[max(0, a0 - context_chars) : a0]
            right = text_a[a1 : a1 + context_chars]
            diffs.append((a0, a1, b0, b1, left, old, new, right))

    cleaned = []
    for a0, a1, b0, b1, left, old, new, right in diffs:
        if normalize_text(old) == normalize_text(new):
            continue
        cleaned.append((a0, a1, b0, b1, left, old, new, right))
    return cleaned


def make_anchor(unit: TextUnit, left_ctx: str, right_ctx: str) -> str:
    h = hashlib.sha1(
        (unit.unit_id + "|" + left_ctx + "|" + right_ctx).encode("utf-8")
    ).hexdigest()[:10]
    return f"{unit.unit_id}:{h}"


def collect_span_diffs(
    aligned_pairs: list[tuple[Optional[TextUnit], Optional[TextUnit]]],
    context_chars: int = 30,
) -> list[SpanDiff]:
    diffs: list[SpanDiff] = []
    for ua, ub in aligned_pairs:
        if not ua or not ub:
            continue
        if ua.text == ub.text:
            continue
        for a0, a1, b0, b1, left, old, new, right in char_span_diffs(
            ua.text, ub.text, context_chars=context_chars
        ):
            diffs.append(SpanDiff(
                unit_id_a=ua.unit_id,
                unit_id_b=ub.unit_id,
                unit_text_a=ua.text,
                unit_text_b=ub.text,
                left_context=left,
                old_span=old,
                new_span=new,
                right_context=right,
                char_start_a=a0,
                char_end_a=a1,
                char_start_b=b0,
                char_end_b=b1,
                anchor=make_anchor(ua, left, right),
            ))
    return diffs


# ---------------------------------------------------------------------------
# Step 4: LLM payload and response schema
# ---------------------------------------------------------------------------

FIELD_TYPES = (
    "person",
    "date",
    "address",
    "county",
    "state",
    "vehicle",
    "license_plate",
    "phone",
    "firm",
    "generic_string",
)


def build_diff_payload(
    span_diffs: list[SpanDiff],
    document_type: str = "SummonsAndComplaint",
) -> dict[str, Any]:
    """Compact payload for LLM: document_type + list of diffs with anchor, context, old/new."""
    return {
        "document_type": document_type,
        "diffs": [
            {
                "anchor": d.anchor,
                "left_context": d.left_context,
                "old_span": d.old_span,
                "new_span": d.new_span,
                "right_context": d.right_context,
            }
            for d in span_diffs
        ],
    }


def build_classify_prompt(diff_json: str) -> str:
    return """You are labeling dynamic placeholders in a legal template.

Given pairs of differing spans from two versions of the SAME template, assign a semantic placeholder field_name for each diff.

Rules:
- Use UPPER_SNAKE_CASE for field_name.
- field_type must be one of: """ + ", ".join(FIELD_TYPES) + """.
- If the span clearly indicates a plaintiff/defendant name, use PLAINTIFF_NAME / DEFENDANT_NAME.
- If it's a venue statement, use VENUE_BASIS or VENUE_ADDRESS as appropriate.
- confidence is 0.0 to 1.0.
- Return JSON ONLY with this exact structure (no markdown, no extra text):

{"mappings": [{"anchor": "<anchor>", "field_name": "FIELD_NAME", "field_type": "person", "confidence": 0.95, "notes": ""}, ...]}

Input diffs:
""" + diff_json


def parse_llm_mappings(response: str) -> dict[str, dict[str, Any]]:
    """Parse LLM response to anchor -> {field_name, field_type, confidence, notes}."""
    text = (response or "").strip()
    if "```" in text:
        parts = text.split("```")
        for p in parts:
            p = p.strip()
            if p.startswith("json"):
                p = p[4:].strip()
            if p.startswith("{"):
                text = p
                break
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        return {}
    mappings = data.get("mappings") or data.get("mapping") or []
    if not isinstance(mappings, list):
        return {}
    result = {}
    for m in mappings:
        if not isinstance(m, dict):
            continue
        anchor = m.get("anchor") or m.get("anchor_id")
        if not anchor:
            continue
        result[str(anchor)] = {
            "field_name": (m.get("field_name") or m.get("field") or "CUSTOM_FIELD").strip().upper().replace(" ", "_"),
            "field_type": m.get("field_type") or "generic_string",
            "confidence": float(m.get("confidence", 0.0)) if m.get("confidence") is not None else 0.0,
            "notes": (m.get("notes") or m.get("note") or "").strip(),
        }
    return result


# ---------------------------------------------------------------------------
# Step 5: Apply placeholders (replace from end to start per unit)
# ---------------------------------------------------------------------------

# Max span length to placeholderize (long spans are likely full sentences, not a single field)
MAX_PLACEHOLDER_SPAN_LENGTH = 60
# Reject span if it contains multiple sentence-ending punctuation (likely boilerplate)
MAX_SENTENCE_PUNCT_IN_SPAN = 1


def _trim_span_by_context(d: SpanDiff, field_name: str) -> tuple[int, int]:
    """
    Narrow the span so we don't placeholderize static prefix/suffix (e.g. "COUNTY OF "
    or "Plaintiff,"). Returns (start, end) in same unit coordinates.
    """
    start, end = d.char_start_a, d.char_end_a
    old = d.old_span.strip()
    left = (d.left_context or "").rstrip()
    # County: replace only the county name, not "COUNTY OF "
    if "COUNTY" in field_name.upper() and old.upper().startswith("COUNTY OF "):
        start = d.char_start_a + len("COUNTY OF ")
    # Plaintiff: replace only the name, not "Plaintiff," or "Plaintiff "
    elif field_name == "PLAINTIFF_NAME":
        if old.startswith("Plaintiff,"):
            start = d.char_start_a + len("Plaintiff,")
        elif old.startswith("Plaintiff "):
            start = d.char_start_a + len("Plaintiff ")
    # Defendant: same
    elif field_name == "DEFENDANT_NAME":
        if old.startswith("Defendant,"):
            start = d.char_start_a + len("Defendant,")
        elif old.startswith("Defendant "):
            start = d.char_start_a + len("Defendant ")
    if start >= end:
        return (d.char_start_a, d.char_end_a)
    return (start, end)


def _should_reject_span(old_span: str, max_length: int = MAX_PLACEHOLDER_SPAN_LENGTH) -> bool:
    """Reject spans that are too long or look like full sentences (multiple commas/periods)."""
    if len(old_span) > max_length:
        return True
    punct_count = sum(1 for c in old_span if c in ".,!?;:")
    if punct_count > MAX_SENTENCE_PUNCT_IN_SPAN:
        return True
    return False


def merge_adjacent_span_diffs_per_unit(
    span_diffs: list[SpanDiff],
    llm_map: dict[str, dict],
    gap_chars: int = 3,
) -> dict[str, list[tuple[int, int, str]]]:
    """
    Merge adjacent or overlapping span diffs within each unit so we replace one
    contiguous span per logical field (avoids KINGSKINGSKINGS / placeholder stacking).
    Skips any diff where unit_text_a at (char_start_a, char_end_a) != old_span (offset guard).
    Returns: unit_id_a -> [(char_start_a, char_end_a, field_name), ...] (non-overlapping, sorted).
    """
    by_unit: dict[str, list[tuple[int, int, str]]] = {}
    for d in span_diffs:
        if d.anchor not in llm_map:
            continue
        # Safety: offset mismatch guard — only use diff if span still matches recorded old_span
        span_text = d.unit_text_a[d.char_start_a : d.char_end_a]
        if span_text != d.old_span and normalize_text(span_text) != normalize_text(d.old_span):
            continue
        # Guard: never placeholderize long spans or full-sentence-like text
        if _should_reject_span(d.old_span):
            continue
        field_name = llm_map[d.anchor].get("field_name", "CUSTOM_FIELD")
        start, end = _trim_span_by_context(d, field_name)
        by_unit.setdefault(d.unit_id_a, []).append((start, end, field_name))

    merged: dict[str, list[tuple[int, int, str]]] = {}
    for unit_id, spans in by_unit.items():
        spans = sorted(spans, key=lambda x: (x[0], x[1]))
        if not spans:
            continue
        run: list[tuple[int, int, str]] = [spans[0]]
        for start, end, field_name in spans[1:]:
            prev_start, prev_end, prev_name = run[-1]
            if start <= prev_end + gap_chars:
                # Overlapping or adjacent: merge (single placeholder for whole span)
                new_end = max(prev_end, end)
                run[-1] = (prev_start, new_end, prev_name)
            else:
                run.append((start, end, field_name))
        merged[unit_id] = run
    return merged


def apply_placeholders_to_units(
    units: list[TextUnit],
    span_diffs: list[SpanDiff],
    llm_map: dict[str, dict],
) -> list[TextUnit]:
    """Return new TextUnits with {{FIELD_NAME}} replacing spans; replace from end to start per unit."""
    by_unit: dict[str, list[SpanDiff]] = {}
    for d in span_diffs:
        if d.anchor in llm_map:
            by_unit.setdefault(d.unit_id_a, []).append(d)

    new_units = []
    for u in units:
        text = u.text
        if u.unit_id in by_unit:
            diffs = sorted(by_unit[u.unit_id], key=lambda x: x.char_start_a, reverse=True)
            for d in diffs:
                field_name = llm_map[d.anchor].get("field_name", "CUSTOM_FIELD")
                text = text[: d.char_start_a] + "{{" + field_name + "}}" + text[d.char_end_a :]
        new_units.append(TextUnit(
            unit_id=u.unit_id,
            text=text,
            raw_text=u.raw_text,
            kind=u.kind,
            path=u.path,
            para_ref=u.para_ref,
        ))
    return new_units


def apply_placeholders_to_docx(
    doc: Document,
    units: list[TextUnit],
    span_diffs: list[SpanDiff],
    llm_map: dict[str, dict],
) -> None:
    """
    Apply placeholder replacements in-place to the document.
    Merges adjacent/overlapping span diffs per unit first to avoid placeholder stacking
    (e.g. KINGSKINGSKINGS). Builds segments from merged spans only; replaces using
    original unit text so offsets are valid. Safety: skip span if text at offset doesn't match.
    """
    from utils.template_builder import _rebuild_paragraph_with_placeholders

    merged_by_unit = merge_adjacent_span_diffs_per_unit(span_diffs, llm_map, gap_chars=3)

    for u in units:
        if u.unit_id not in merged_by_unit or not u.para_ref:
            continue
        para = u.para_ref
        full_text = u.text
        merged_spans = merged_by_unit[u.unit_id]
        # Build segments (start, end, key) for _rebuild_paragraph_with_placeholders; no overlaps
        segments: list[tuple[int, int, Optional[str]]] = []
        pos = 0
        for start, end, field_name in merged_spans:
            # Safety: clamp to unit text bounds; skip if empty or out of range
            start = max(0, min(start, len(full_text)))
            end = max(start, min(end, len(full_text)))
            if start >= end:
                continue
            # Guard: ensure we are not replacing more than the paragraph (sanity)
            if end - start > len(full_text) * 0.95:
                continue
            if pos < start:
                segments.append((pos, start, None))
            segments.append((start, end, field_name))
            pos = end
        if pos < len(full_text):
            segments.append((pos, len(full_text), None))
        _rebuild_paragraph_with_placeholders(para, full_text, segments)


# ---------------------------------------------------------------------------
# Hardening: filter boilerplate, optional merge adjacent spans
# ---------------------------------------------------------------------------


def filter_boilerplate_diffs(
    span_diffs: list[SpanDiff],
    max_span_length: int = 120,
    min_similarity_ratio: float = 0.3,
) -> list[SpanDiff]:
    """Drop diffs where both old and new are long and very different (likely structure drift)."""
    kept = []
    for d in span_diffs:
        if len(d.old_span) > max_span_length and len(d.new_span) > max_span_length:
            r = difflib.SequenceMatcher(None, d.old_span, d.new_span, autojunk=False).ratio()
            if r < min_similarity_ratio:
                continue
        kept.append(d)
    return kept


# ---------------------------------------------------------------------------
# Entry point: infer placeholders from two DOCX
# ---------------------------------------------------------------------------


def infer_placeholders_from_two_docx(
    sample_a_path: str | Path,
    sample_b_path: str | Path,
    document_type: str = "SummonsAndComplaint",
    *,
    use_llm: bool = False,
    llm_callable: Optional[Any] = None,
    filter_boilerplate: bool = True,
    context_chars: int = 30,
) -> dict[str, Any]:
    """
    Run deterministic diff + optional LLM classification. Returns blueprint dict:

    - span_diffs: list of serialized SpanDiff dicts
    - llm_map: anchor -> {field_name, field_type, confidence, notes}
    - placeholder_keys: list of unique field names
    - units_a, span_diffs_raw: for apply_placeholders_to_docx (if generating template)

    Does NOT write template.docx; caller uses apply_placeholders_to_docx then doc.save().
    """
    path_a = Path(sample_a_path)
    path_b = Path(sample_b_path)
    doc_a = Document(path_a)
    doc_b = Document(path_b)

    units_a = extract_units(doc_a, str(path_a))
    units_b = extract_units(doc_b, str(path_b))

    aligned = align_units(units_a, units_b)
    span_diffs = collect_span_diffs(aligned, context_chars=context_chars)

    if filter_boilerplate:
        span_diffs = filter_boilerplate_diffs(span_diffs)

    llm_map: dict[str, dict] = {}
    if use_llm and callable(llm_callable):
        payload = build_diff_payload(span_diffs, document_type=document_type)
        diff_json = json.dumps(payload, indent=2, ensure_ascii=False)
        prompt = build_classify_prompt(diff_json)
        try:
            response = llm_callable(prompt)
            llm_map = parse_llm_mappings(response)
        except Exception:
            pass

    for d in span_diffs:
        if d.anchor in llm_map:
            continue
        llm_map[d.anchor] = _heuristic_classify(d)

    placeholder_keys = sorted(set(m.get("field_name", "CUSTOM_FIELD") for m in llm_map.values()))

    return {
        "span_diffs": [
            {
                "unit_id_a": d.unit_id_a,
                "unit_id_b": d.unit_id_b,
                "left_context": d.left_context,
                "old_span": d.old_span,
                "new_span": d.new_span,
                "right_context": d.right_context,
                "char_start_a": d.char_start_a,
                "char_end_a": d.char_end_a,
                "anchor": d.anchor,
            }
            for d in span_diffs
        ],
        "span_diffs_raw": span_diffs,
        "units_a": units_a,
        "llm_map": llm_map,
        "placeholder_keys": placeholder_keys,
        "document_type": document_type,
        "doc_a": doc_a,
    }


def _heuristic_classify(d: SpanDiff) -> dict[str, Any]:
    """Heuristic classification when LLM is not used or fails."""
    left = (d.left_context or "").lower()
    right = (d.right_context or "").lower()
    old = (d.old_span or "").lower()
    new = (d.new_span or "").lower()
    combined = left + " " + right
    if "plaintiff" in combined and ("against" in left or "plaintiff" in left):
        return {"field_name": "PLAINTIFF_NAME", "field_type": "person", "confidence": 0.8, "notes": ""}
    if "defendant" in combined:
        return {"field_name": "DEFENDANT_NAME", "field_type": "person", "confidence": 0.8, "notes": ""}
    if "index" in combined or "index no" in combined:
        return {"field_name": "INDEX_NO", "field_type": "generic_string", "confidence": 0.85, "notes": ""}
    if "date filed" in combined or "filed" in combined:
        return {"field_name": "DATE_FILED", "field_type": "date", "confidence": 0.85, "notes": ""}
    if "venue" in combined or "basis" in combined:
        return {"field_name": "VENUE_BASIS", "field_type": "address", "confidence": 0.75, "notes": ""}
    if "county" in combined:
        return {"field_name": "COUNTY_NAME", "field_type": "county", "confidence": 0.8, "notes": ""}
    if "vehicle" in combined or "license" in combined or "plate" in combined:
        if re.search(r"\b(19|20)\d{2}\b", d.old_span + d.new_span):
            return {"field_name": "VEHICLE_YEAR", "field_type": "vehicle", "confidence": 0.8, "notes": ""}
        if any(x in old + new for x in ["lexus", "toyota", "honda", "ford"]):
            return {"field_name": "VEHICLE_MAKE", "field_type": "vehicle", "confidence": 0.8, "notes": ""}
        return {"field_name": "LICENSE_PLATE", "field_type": "license_plate", "confidence": 0.75, "notes": ""}
    if "address" in combined or "road" in combined or "avenue" in combined or "street" in combined:
        return {"field_name": "VENUE_ADDRESS", "field_type": "address", "confidence": 0.75, "notes": ""}
    if re.search(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\s+\d{1,2},?\s+\d{4}\b", d.old_span + d.new_span, re.I):
        return {"field_name": "ACCIDENT_DATE", "field_type": "date", "confidence": 0.8, "notes": ""}
    if "attorney" in combined or "esq" in combined:
        return {"field_name": "SIGNATURE_BLOCK.ATTORNEY", "field_type": "person", "confidence": 0.7, "notes": ""}
    if "firm" in combined or "pllc" in combined or "p.c." in combined:
        return {"field_name": "SIGNATURE_BLOCK.FIRM", "field_type": "firm", "confidence": 0.7, "notes": ""}
    if "phone" in combined or re.search(r"\(\d{3}\)", d.old_span + d.new_span):
        return {"field_name": "SIGNATURE_BLOCK.PHONE", "field_type": "phone", "confidence": 0.7, "notes": ""}
    return {"field_name": "CUSTOM_FIELD", "field_type": "generic_string", "confidence": 0.5, "notes": ""}
